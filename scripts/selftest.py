"""
selftest.py — офлайн-проверка работоспособности парсера БЕЗ обращения к Avito.

Проверяет всю внутреннюю «машинерию»: импорт, фабрику уведомлений, доставку
webhook (на локальный mock-сервер), дедупликацию SQLite + авто-миграцию,
устойчивый парсинг объявлений и атомарную запись Excel.

Запуск (из корня проекта, в активированном venv):
    python scripts/selftest.py

Код возврата 0 — все проверки пройдены, иначе 1.
"""
import json
import os
import sqlite3
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler
from pathlib import Path
from socketserver import TCPServer

# чтобы скрипт работал из любой директории
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

PASSED, FAILED = [], []


def check(name, cond):
    (PASSED if cond else FAILED).append(name)
    print(f"  [{'OK' if cond else 'FAIL'}] {name}")


def main():
    # ── mock-сервер для webhook ──
    received = []

    class H(BaseHTTPRequestHandler):
        def do_POST(self):
            n = int(self.headers.get("Content-Length", 0))
            received.append({
                "secret": self.headers.get("X-Webhook-Secret"),
                "json": json.loads(self.rfile.read(n).decode()),
            })
            self.send_response(200); self.end_headers(); self.wfile.write(b'{"ok":true}')

        def log_message(self, *a):
            pass

    httpd = TCPServer(("127.0.0.1", 0), H)
    port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{port}/avito"

    print("1) Импорт проекта")
    import parser_cls
    from dto import AvitoConfig
    from models import Item
    from integrations.notifications.factory import build_notifier
    from db_service import SQLiteDBHandler
    from parser.export.excel import ExcelStorage
    check("import parser_cls + модули", True)

    print("2) Фабрика уведомлений (webhook)")
    cfg = AvitoConfig(urls=["https://avito.ru/searchA"], webhook_url=url, webhook_secret="s3cr3t")
    notifier = build_notifier(cfg)
    check("build_notifier -> CompositeNotifier[WebhookNotifier]",
          type(notifier).__name__ == "CompositeNotifier"
          and any(type(n).__name__ == "WebhookNotifier" for n in notifier.notifiers))

    print("3) Доставка webhook с source-тегом")
    ad = Item(id=987654321, urlPath="krasnodar/telefony/iphone_16", title="iPhone 16 тест", sellerId="seller-xyz")
    notifier.notify_many(ads=[ad], source="https://avito.ru/searchA")
    time.sleep(0.3)
    ok = bool(received) and received[0]["secret"] == "s3cr3t" \
        and received[0]["json"].get("source") == "https://avito.ru/searchA" \
        and received[0]["json"].get("id") == 987654321
    check("webhook доставлен (source + secret + payload)", ok)

    print("4) SQLite: миграция старой БД + дедуп + WAL")
    tmp = tempfile.mkdtemp()
    dbp = os.path.join(tmp, "old.db")
    c = sqlite3.connect(dbp)
    c.execute("CREATE TABLE viewed (id INTEGER, price INTEGER)")  # старая схема без PK
    c.executemany("INSERT INTO viewed VALUES (?,?)", [(1, 100), (1, 100), (1, 100), (2, 200)])
    c.commit(); c.close()
    SQLiteDBHandler._instance = None
    db = SQLiteDBHandler(db_name=dbp)
    c = sqlite3.connect(dbp)
    rows = c.execute("SELECT COUNT(*) FROM viewed").fetchone()[0]
    pk = {r[1] for r in c.execute("PRAGMA table_info(viewed)").fetchall() if r[5] > 0}
    wal = c.execute("PRAGMA journal_mode").fetchone()[0].lower()
    c.close()
    check("миграция: 4 строки -> 2 (дедуп)", rows == 2)
    check("PRIMARY KEY (id, price)", pk == {"id", "price"})
    check("journal_mode = WAL", wal == "wal")
    check("record_exists по (id, price)", db.record_exists(1, 100) and not db.record_exists(1, 999))

    print("5) Устойчивый парсинг (битое объявление не роняет страницу)")
    good = {"id": 111, "title": "ok", "urlPath": "a/b"}
    bad = {"id": 222, "priceDetailed": {"enabled": True}}  # неполный PriceDetailed -> ошибка валидации
    ads = parser_cls.AvitoParse._parse_items_tolerant({"items": [good, bad, good]})
    check("из 3 (1 битое) распарсено 2", len(ads) == 2)

    print("6) Атомарная запись Excel")
    fp = Path(tmp) / "avito.xlsx"
    st = ExcelStorage(fp)

    class P:
        value = 75000

    def mkad(i):
        a = Item(id=i, title=f"ad{i}", urlPath=f"x/{i}", sortTimeStamp=1700000000000, images=[])
        a.priceDetailed = P()
        return a

    st.save([mkad(1)])
    st.save([mkad(2), mkad(3)])
    from openpyxl import load_workbook
    n_rows = load_workbook(fp).active.max_row
    leftover = [f for f in os.listdir(tmp) if f.endswith(".tmp")]
    check("строки накапливаются (заголовок + 3)", n_rows == 4)
    check("нет временных .tmp", not leftover)

    httpd.shutdown()

    print("\n" + "=" * 48)
    print(f"ПРОЙДЕНО: {len(PASSED)} | ПРОВАЛЕНО: {len(FAILED)}")
    if FAILED:
        print("Провалены:", ", ".join(FAILED))
        return 1
    print("ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ ✅ — внутренняя логика работает.")
    print("Следующий шаг: тест webhook в свой n8n (scripts/send_test_webhook.py)")
    print("и боевой прогон к Avito (см. docs/HOW_TO_TEST.md).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
